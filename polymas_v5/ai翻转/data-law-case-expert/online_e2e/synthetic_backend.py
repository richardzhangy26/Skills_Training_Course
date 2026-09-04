"""仅供离线状态机回归的 synthetic backend。"""

from __future__ import annotations

from collections.abc import Mapping
from copy import deepcopy
from io import BytesIO
from typing import Any

from docx import Document

from .backends import BackendFailure, Blocker, PrecheckResult, RegressionSnapshot
from .clients import PdsClient
from .config_diff import snapshot_knowledge
from .contracts import ConfigSnapshot, TargetConfig


def _thaw(value):
    if isinstance(value, Mapping):
        return {str(key): _thaw(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_thaw(item) for item in value]
    return value


class SyntheticRegressionBackend:
    environment = "synthetic"

    def __init__(
        self,
        target: TargetConfig,
        *,
        blockers: tuple[Blocker, ...] = (),
        fail_at: str | None = None,
        external_change_before_rollback: bool = False,
        malformed_teacher_receipt: bool = False,
        malformed_publish_receipt: bool = False,
        external_knowledge_change_before_cleanup: bool = False,
        corrupt_student: str | None = None,
        publish_assistant_id: str = "synthetic-isolated-assistant",
        fake_cleanup_receipt: bool = False,
    ):
        self.target = target
        self.blockers = blockers
        self.fail_at = fail_at
        self.external_change_before_rollback = external_change_before_rollback
        self.malformed_teacher_receipt = malformed_teacher_receipt
        self.malformed_publish_receipt = malformed_publish_receipt
        self.external_knowledge_change_before_cleanup = external_knowledge_change_before_cleanup
        self.corrupt_student = corrupt_student
        self.publish_assistant_id = publish_assistant_id
        self.fake_cleanup_receipt = fake_cleanup_receipt
        self.precheck_assistant_id = "synthetic-isolated-assistant"
        self.relationship_version = "synthetic-v1"
        self.write_count = 0
        self.restore_config_calls = 0
        self.student_calls: list[str] = []
        self.temporary_cases: set[str] = set()
        self.knowledge_version = "knowledge-v1"
        self.knowledge_cas_checks = 0
        self._knowledge_reads_after_sync = 0
        self._knowledge_synced = False
        self._knowledge_content = b"synthetic baseline knowledge"
        self._full_config = self._make_full_config("synthetic baseline Agent.md")

    def _make_full_config(self, content: str) -> dict[str, Any]:
        return {
            "basicInfo": {
                "nid": self.target.expert_nid,
                "appName": self.target.expert_name,
                "appType": "EXPERT",
                "isPublish": 1,
            },
            "personaMd": None,
            "soulMd": None,
            "agentMd": None,
            "planMd": None,
            "expertMd": {
                "templateNid": "eu5FW3sdWS",
                "customContent": content,
                "rawContent": "synthetic raw content",
            },
            "skillInfoList": [
                {
                    "skillNid": self.target.online_skill_nids[name],
                    "name": name,
                    "enabled": True,
                    "version": "synthetic",
                    "bindingSource": self.target.online_skill_binding_sources[name],
                    "permission": "TEACHER",
                }
                for name in self.target.expected_skill_order
            ],
            "subAgentVOS": [],
            "mcpInfoList": [],
            "datasets": None,
            "generalSetting": {},
            "agentLlmModelConfig": {},
            "extInfo": {"recommendedQuestions": [], "expertiseAreas": [], "teamScene": None},
        }

    @property
    def agent_content(self) -> str:
        return self._full_config["expertMd"]["customContent"]

    def _config_snapshot(self) -> ConfigSnapshot:
        return PdsClient.snapshot_from_config(self.target.expert_nid, self._full_config, [])

    def precheck(self, target: TargetConfig) -> PrecheckResult:
        return PrecheckResult(
            blockers=self.blockers,
            relationship_version=self.relationship_version,
            assistant_nid=self.precheck_assistant_id,
        )

    def snapshot(self, target: TargetConfig) -> RegressionSnapshot:
        return RegressionSnapshot(
            config=self._config_snapshot(),
            knowledge=snapshot_knowledge(
                self.knowledge_version,
                self._knowledge_content,
                source="synthetic-content-snapshot",
            ),
        )

    def _knowledge_snapshot(self):
        return snapshot_knowledge(
            self.knowledge_version,
            self._knowledge_content,
            source="synthetic-content-snapshot",
        )

    def current_knowledge_snapshot(self):
        self.knowledge_cas_checks += 1
        if (
            self._knowledge_synced
            and self.external_knowledge_change_before_cleanup
            and self._knowledge_reads_after_sync >= 1
            and self.knowledge_version != "external-knowledge-version"
        ):
            self.knowledge_version = "external-knowledge-version"
            self._knowledge_content = b"external knowledge content"
        if self._knowledge_synced:
            self._knowledge_reads_after_sync += 1
        return self._knowledge_snapshot()

    def publish(self, target: TargetConfig, desired: Mapping[str, Any], run_id: str):
        if self.fail_at == "publish":
            raise BackendFailure("SYNTHETIC_PUBLISH_FAILED")
        self.write_count += 1
        self._full_config = deepcopy(dict(desired))
        if self.malformed_publish_receipt:
            return {"answer": "published"}
        digest = self._config_snapshot().digest
        return {"published": True, "owned_digest": digest,
                "assistantId": self.publish_assistant_id}

    def current_config_digest(self, target: TargetConfig) -> str:
        return self._config_snapshot().digest

    def run_student(self, scenario, assistant_nid: str):
        self.write_count += 1
        self.student_calls.append(scenario.scenario_id)
        if self.fail_at == f"student:{scenario.scenario_id}":
            if self.external_change_before_rollback:
                self._full_config["expertMd"]["customContent"] = "external concurrent content"
            raise BackendFailure("SYNTHETIC_STUDENT_FAILED", scenario.scenario_id)
        common = {
            "passed": True,
            "scenario_id": scenario.scenario_id,
            "assistantId": assistant_nid,
            "conversationId": "synthetic-conversation",
            "messageId": f"synthetic-message-{scenario.scenario_id}",
            "planId": f"synthetic-plan-{scenario.scenario_id}",
            "traceId": f"synthetic-trace-{scenario.scenario_id}",
        }
        receipts = {
            "exact-statute": {
                "outcome": "answered", "answer": "DLCL-0001 statute evidence and reflection question",
                "caseIds": ["DLCL-0001"],
                "evidence": {"statutes": ["synthetic-statute"], "reflectionQuestion": True},
            },
            "detailed-explanation": {
                "outcome": "answered", "answer": "DLCL-0001 facts and dispute with a follow-up question",
                "caseIds": ["DLCL-0001"],
                "evidence": {"caseSummary": True, "dispute": True, "followUpQuestion": True},
            },
            "follow-up-question": {
                "outcome": "answered", "answer": "The dispute focus concerns responsibility.",
                "caseIds": ["DLCL-0001"], "evidence": {"focus": "dispute"},
            },
            "ambiguous-candidates": {
                "outcome": "awaiting_selection", "answer": "Choose one candidate.",
                "candidates": [{"caseId": "DLCL-0001"}, {"caseId": "DLCL-0002"}],
            },
            "unknown-case-no-fabrication": {
                "outcome": "not_found", "answer": "No matching case.", "caseIds": [],
                "evidence": {"fabricatedFacts": False, "inventedCitation": False},
            },
            "student-write-denied": {
                "outcome": "denied", "answer": "Write denied.", "writePerformed": False,
                "reasonCode": "ROLE_NOT_AUTHORIZED",
            },
        }
        receipt = {**common, **receipts[scenario.scenario_id]}
        if self.corrupt_student == scenario.scenario_id:
            receipt["outcome"] = "answered"
            receipt["candidates"] = []
            receipt["passed"] = True
        return receipt

    def upload_teacher_fixture(self, payload: bytes, *, scene: str, case_ids: tuple[str, ...], run_id: str):
        self.write_count += 1
        try:
            document = Document(BytesIO(payload))
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            text += "\n" + "\n".join(
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            )
        except Exception:
            raise BackendFailure("FIXTURE_INVALID") from None
        required = (*case_ids, run_id, "FICTIONAL TEST CASES", "NO REAL PII")
        if (
            scene != "自动化测试"
            or len(case_ids) != 2
            or len(set(case_ids)) != 2
            or any(text.count(value) != 1 for value in case_ids)
            or any(value not in text for value in required[2:])
        ):
            raise BackendFailure("FIXTURE_INVALID")
        if self.malformed_teacher_receipt:
            return {"answer": "上传成功"}
        if self.fail_at == "teacher:upload":
            raise BackendFailure("SYNTHETIC_UPLOAD_FAILED")
        return {"accepted": True, "uploadId": f"upload-{run_id}", "scene": scene}

    def confirm_teacher_change(self, upload_id: str):
        self.write_count += 1
        if self.fail_at == "teacher:confirm":
            raise BackendFailure("SYNTHETIC_CONFIRM_FAILED")
        return {"confirmed": True, "changeId": f"change-{upload_id}"}

    def sync_teacher_change(self, change_id: str):
        self.write_count += 1
        if self.fail_at == "teacher:sync":
            raise BackendFailure("SYNTHETIC_SYNC_FAILED")
        self.temporary_cases.update(
            item for item in getattr(self, "_pending_case_ids", ())
        )
        self.knowledge_version = "knowledge-temporary"
        self._knowledge_content = b"synthetic temporary knowledge"
        self._knowledge_synced = True
        self._knowledge_reads_after_sync = 0
        owned = self._knowledge_snapshot()
        return {
            "htmlUpdated": True,
            "knowledgeUpdated": True,
            "knowledgeVersion": owned.version,
            "knowledgeDigest": owned.digest,
        }

    def read_case(self, case_id: str):
        if case_id not in self.temporary_cases:
            return None
        return {"caseId": case_id, "scene": "自动化测试"}

    def cleanup_teacher_cases(self, case_ids: tuple[str, ...], run_id: str):
        self.write_count += 1
        if self.fail_at == "cleanup:cases":
            raise BackendFailure("SYNTHETIC_CLEANUP_FAILED")
        deleted = [case_id for case_id in case_ids if case_id in self.temporary_cases]
        if not self.fake_cleanup_receipt:
            self.temporary_cases.difference_update(deleted)
        return {"cleaned": True, "deletedIds": deleted, "ownedRunId": run_id}

    def restore_knowledge(self, snapshot, *, owned_version: str, owned_digest: str):
        self.write_count += 1
        if self.fail_at == "cleanup:knowledge":
            raise BackendFailure("SYNTHETIC_KNOWLEDGE_RESTORE_FAILED")
        current = self._knowledge_snapshot()
        if current.version != owned_version or current.digest != owned_digest:
            raise BackendFailure("EXTERNAL_CONCURRENT_CHANGE", "knowledge_not_restored")
        self.knowledge_version = snapshot.version
        self._knowledge_content = b"synthetic baseline knowledge"
        self._knowledge_synced = False
        actual = snapshot_knowledge(self.knowledge_version, self._knowledge_content,
                                    source="synthetic-content-snapshot")
        return {"restored": actual.digest == snapshot.digest,
                "knowledgeVersion": actual.version, "knowledgeDigest": actual.digest}

    def verify_cases_absent(self, case_ids: tuple[str, ...]) -> bool:
        return not self.temporary_cases.intersection(case_ids)

    def restore_config(self, snapshot: ConfigSnapshot, *, owned_digest: str):
        self.restore_config_calls += 1
        self.write_count += 1
        if self.current_config_digest(self.target) != owned_digest:
            raise BackendFailure("EXTERNAL_CONCURRENT_CHANGE")
        self._full_config = _thaw(snapshot.normalized["full_config"])
        return {"restored": self.current_config_digest(self.target) == snapshot.digest}

    def remember_pending_cases(self, case_ids: tuple[str, ...]) -> None:
        self._pending_case_ids = case_ids
